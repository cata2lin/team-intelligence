# /// script
# requires-python = ">=3.11"
# dependencies = ["xlrd==1.2.0", "openpyxl", "xlwt"]
# ///
"""Teste pentru deconturi_ramburs. Rulează: `uv run test_deconturi_ramburs.py` (nu cere mail, nu cere rețea).

Fișierele de probă se construiesc aici, cu exact capcanele întâlnite pe datele reale din aug–oct 2026 (valorile sunt fictive):
AWB-uri venite din Excel ca float, coloana de detaliu Sameday pe A DOUA foaie, și fișierele FAN care
seamănă a decont dar nu sunt.
"""
import os, sys, tempfile
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from deconturi_ramburs import parseaza_fisier, parseaza_director, CAMPURI


def _dpd(cale):
    import xlwt
    wb = xlwt.Workbook(); ws = wb.add_sheet("Sheet1")
    for j, v in enumerate(["Line", "Waybill", "Pickup date", "Delivery date", "Sender", "Recipient",
                           "Note(Client)", "Reference 1", "Reference 2", "OrderID", "Amount"]):
        ws.write(0, j, v)
    # AWB-ul vine ca NUMĂR din Excel -> „10000000001.0"; suma la fel
    ws.write(1, 0, 1.0); ws.write(1, 1, 10000000001.0); ws.write(1, 3, "21.10.2026")
    ws.write(1, 7, "BON000001 "); ws.write(1, 10, 169.0)
    ws.write(2, 0, 2.0); ws.write(2, 1, 10000000002.0); ws.write(2, 3, "21.10.2026")
    ws.write(2, 7, "EST000001 "); ws.write(2, 10, 300.0)
    ws.write(3, 0, 3.0)                       # rând fără AWB: se sare, nu crapă
    wb.save(cale)


def _sameday(cale):
    import openpyxl
    wb = openpyxl.Workbook(); c = wb.active; c.title = "client"
    c.append(["Nume client / companie", "Brand", "Banca", "IBAN", "Suma totala", "Perioada export"])
    c.append(["ARONA SRL", "ARONA SRL", "Wise", "RO41…", 839.18, "15 Oct 2026 - 18 Oct 2026"])
    e = wb.create_sheet("expeditii")
    e.append(["AWB", "Nume destinatar", "Judet destinatar", "Oras destinatar", "Adresa destinatar",
              "Adresa ridicare", "Referinta client", "Suma ramburs"])
    e.append(["1ONTEST000000001", "Client Test", "Judet", "Localitate", "Str. X", "Depozit", "TEST000001_A_1", 114])
    wb.save(cale)


def _fan(cale):
    import openpyxl
    wb = openpyxl.Workbook(); ws = wb.active
    ws.append(["AWB", "Client", "Localitate", "Judet", "Adresa", "Continut", "ValoareDeclarata", "Ramburs"])
    ws.append(["7000000000001", "Client Test", "Localitate", "Judet", "Str. Y", "", "0.00", "150"])
    wb.save(cale)


def test_dpd():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "COD_0000000_2026.10.21_000000000000000001_000001.xls")
        _dpd(p)
        r = parseaza_fisier(p)
        assert len(r) == 2, f"asteptam 2 randuri, am primit {len(r)}"
        a = r[0]
        assert a["awb"] == "10000000001", f"AWB-ul a ramas cu zecimale: {a['awb']}"
        assert a["suma"] == 169.0 and a["moneda"] == "RON"
        assert a["data_virarii"] == "2026-10-21", a["data_virarii"]
        assert a["data_livrarii"] == "2026-10-21", a["data_livrarii"]
        assert a["referinta_client"] == "BON000001", repr(a["referinta_client"])
        assert a["curier"] == "DPD" and a["tip"] == "decont"
        assert set(a) == set(CAMPURI), "campurile s-au schimbat fara sa se schimbe contractul"
        print("OK: DPD — AWB fara zecimale, data din nume, referinta curatata")


def test_sameday_detaliul_e_pe_a_doua_foaie():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "2026-10-19-arona-srl-ramburs-de-transferat-ron.xlsx")
        _sameday(p)
        r = parseaza_fisier(p)
        assert len(r) == 1, f"detaliul Sameday nu a fost citit de pe foaia `expeditii`: {len(r)} randuri"
        assert r[0]["awb"] == "1ONTEST000000001" and r[0]["suma"] == 114.0
        assert r[0]["curier"] == "SAMEDAY" and r[0]["data_virarii"] == "2026-10-19"
        print("OK: Sameday — citeste foaia `expeditii`, nu prima (care are doar totalul)")


def test_fan_nu_e_decont():
    with tempfile.TemporaryDirectory() as d:
        p = os.path.join(d, "Expeditii_FAN_ARONA_SRL_Brasov_15.09.2026.xlsx")
        _fan(p)
        r = parseaza_fisier(p)
        assert len(r) == 1 and r[0]["tip"] == "expeditii", \
            "fisierele FAN sunt liste de expeditii; daca intra ca `decont` se dubleaza banii in reconciliere"
        assert r[0]["data_virarii"] is None, "o lista de expeditii nu are data de virare"
        print("OK: FAN — marcat `expeditii`, nu intra in reconciliere ca decont")


def test_fisier_strain_si_lot():
    with tempfile.TemporaryDirectory() as d:
        _dpd(os.path.join(d, "COD_0000000_2026.10.21_000000000000000001_000001.xls"))
        _sameday(os.path.join(d, "2026-10-19-arona-srl-ramburs-de-transferat-ron.xlsx"))
        open(os.path.join(d, "altceva.xlsx"), "wb").write(b"nu e un xlsx valid")
        open(os.path.join(d, "_index.json"), "w").write("{}")
        r = parseaza_director(d)      # fisierul stricat NU trebuie sa opreasca lotul
        assert len(r) == 3, f"asteptam 3 randuri din lot, am primit {len(r)}"
        assert {x["curier"] for x in r} == {"DPD", "SAMEDAY"}
        print("OK: lotul continua peste un fisier stricat; `_index.json` e ignorat")


if __name__ == "__main__":
    test_dpd(); test_sameday_detaliul_e_pe_a_doua_foaie()
    test_fan_nu_e_decont(); test_fisier_strain_si_lot()
    print("TOATE TRECUT")
